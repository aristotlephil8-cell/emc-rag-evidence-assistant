from __future__ import annotations

import zipfile
from io import BytesIO

import pytest
from docx import Document

from app.ingestion import (
    PARSER_BUNDLE_VERSION,
    IngestionError,
    IngestionErrorCode,
    ParserService,
)
from app.ingestion.pdf_plain import PDF_MAX_BYTES
from app.schemas.ingestion import BlockType, FileType


def _pdf_bytes(page_texts: list[str], *, vector_grid: bool = False) -> bytes:
    """Build a small deterministic Type-1-font PDF without test-only dependencies."""

    objects: dict[int, bytes] = {}
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    page_ids: list[int] = []
    for index, text in enumerate(page_texts):
        page_id = 4 + index * 2
        stream_id = page_id + 1
        page_ids.append(page_id)
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 14 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
        if vector_grid:
            stream += (
                b"\n72 680 m 300 680 l S 72 650 m 300 650 l S "
                b"72 620 m 300 620 l S 72 620 m 72 680 l S "
                b"180 620 m 180 680 l S 300 620 m 300 680 l S"
            )
        objects[page_id] = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {stream_id} 0 R >>"
        ).encode("ascii")
        objects[stream_id] = (
            f"<< /Length {len(stream)} >>\nstream\n".encode("ascii") + stream + b"\nendstream"
        )
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[2] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode("ascii")

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0] * (max(objects) + 1)
    for object_id in sorted(objects):
        offsets[object_id] = len(output)
        output.extend(f"{object_id} 0 obj\n".encode("ascii"))
        output.extend(objects[object_id])
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(output)


def test_parse_utf8_txt_with_heading_and_locator() -> None:
    parsed = ParserService.parse(
        "guide.txt",
        b"# EMC Guide\n\nBond the shield at the enclosure entry.",
    )

    assert parsed.file_type == FileType.TXT
    assert parsed.parser_version == "txt-1.0.0"
    assert parsed.sections[0].block_type == BlockType.HEADING
    assert parsed.sections[1].section_path == ("EMC Guide",)
    assert parsed.sections[1].metadata["line_start"] == 3


def test_txt_rejects_non_utf8() -> None:
    with pytest.raises(IngestionError) as caught:
        ParserService.parse("bad.txt", b"\xff\xfe\xfa")
    assert caught.value.code == IngestionErrorCode.UNSUPPORTED_TEXT_ENCODING


def test_parse_docx_heading_paragraph_and_table() -> None:
    document = Document()
    document.add_heading("Cable Shielding", level=1)
    document.add_paragraph("Terminate the braid with a low-impedance connection.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Frequency"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "30 MHz"
    table.cell(1, 1).text = "40 dB"
    buffer = BytesIO()
    document.save(buffer)

    parsed = ParserService.parse("procedure.docx", buffer.getvalue())

    assert parsed.file_type == FileType.DOCX
    assert [section.block_type for section in parsed.sections[:2]] == [
        BlockType.HEADING,
        BlockType.PARAGRAPH,
    ]
    table_rows = [section for section in parsed.sections if section.block_type == BlockType.TABLE]
    assert [section.table_row for section in table_rows] == [1, 2]
    assert all(section.section_path == ("Cable Shielding",) for section in table_rows)


def test_docx_entry_limit_rejects_zip_bomb_shape() -> None:
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as package:
        package.writestr("[Content_Types].xml", "types")
        package.writestr("word/document.xml", "document")
        for index in range(1023):
            package.writestr(f"word/padding/{index}.xml", "")

    with pytest.raises(IngestionError) as caught:
        ParserService.parse("bomb.docx", buffer.getvalue())
    assert caught.value.code == IngestionErrorCode.RESOURCE_LIMIT_EXCEEDED
    assert "too many entries" in caught.value.safe_message


def test_docx_total_uncompressed_limit_is_checked(monkeypatch: pytest.MonkeyPatch) -> None:
    document = Document()
    document.add_paragraph("EMC")
    buffer = BytesIO()
    document.save(buffer)
    monkeypatch.setattr("app.ingestion.docx.DOCX_MAX_UNCOMPRESSED_BYTES", 1)

    with pytest.raises(IngestionError) as caught:
        ParserService.parse("large.docx", buffer.getvalue())
    assert caught.value.code == IngestionErrorCode.RESOURCE_LIMIT_EXCEEDED
    assert "uncompressed size" in caught.value.safe_message


def test_parse_plain_pdf_with_page_locator() -> None:
    parsed = ParserService.parse(
        "plain.pdf",
        _pdf_bytes(["EMC shielding verification procedure"]),
    )

    assert parsed.file_type == FileType.PDF
    assert parsed.metadata["route"] == "plain"
    assert parsed.sections[0].page_number == 1
    assert parsed.sections[0].bbox is not None
    assert parsed.sections[0].bbox != (0.0, 0.0, 612.0, 792.0)
    assert parsed.sections[0].metadata["bbox_precision"] == "object_line"
    assert "shielding" in parsed.sections[0].text


def test_vector_table_routes_complete_pdf_to_deepdoc() -> None:
    with pytest.raises(IngestionError) as caught:
        ParserService.parse(
            "vector-table.pdf",
            _pdf_bytes(["Frequency Limit"], vector_grid=True),
        )
    assert caught.value.code == IngestionErrorCode.MODEL_ASSET_MISSING


def test_pdf_enforces_20_mib_limit_before_parsing() -> None:
    oversized = b"%PDF-1.4\n" + b"x" * PDF_MAX_BYTES
    with pytest.raises(IngestionError) as caught:
        ParserService.parse("large.pdf", oversized)
    assert caught.value.code == IngestionErrorCode.RESOURCE_LIMIT_EXCEEDED
    assert "20 MiB" in caught.value.safe_message


def test_pdf_enforces_50_page_limit() -> None:
    with pytest.raises(IngestionError) as caught:
        ParserService.parse("many-pages.pdf", _pdf_bytes(["EMC"] * 51))
    assert caught.value.code == IngestionErrorCode.RESOURCE_LIMIT_EXCEEDED
    assert "50 page" in caught.value.safe_message


def test_scanned_pdf_requires_explicit_model_directory() -> None:
    with pytest.raises(IngestionError) as caught:
        ParserService.parse("scan.pdf", _pdf_bytes([""]))
    assert caught.value.code == IngestionErrorCode.MODEL_ASSET_MISSING


def test_unsupported_extension_is_explicit() -> None:
    with pytest.raises(IngestionError) as caught:
        ParserService.parse("notes.md", b"text")
    assert caught.value.code == IngestionErrorCode.UNSUPPORTED_FILE_TYPE


def test_parser_bundle_version_covers_every_parser_without_onnx_import() -> None:
    assert "txt=txt-1.0.0" in PARSER_BUNDLE_VERSION
    assert "docx=docx-1.1.0" in PARSER_BUNDLE_VERSION
    assert "pdf=pdf-plain-1.1.0" in PARSER_BUNDLE_VERSION
    assert "deepdoc=deepdoc-de0e793dc6d7-1.1.0" in PARSER_BUNDLE_VERSION
